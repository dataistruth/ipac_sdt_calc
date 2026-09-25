#!/usr/bin/env python3
"""Generate the uspGetFinalEffectivePercentage optimization review deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "uspGetFinalEffectivePercentage_Optimization_Review.pptx"

W = Inches(13.333)
H = Inches(7.5)

GREEN = RGBColor(134, 188, 37)
DARK_GREEN = RGBColor(44, 83, 29)
NAVY = RGBColor(20, 39, 63)
BLUE = RGBColor(37, 99, 160)
LIGHT_BLUE = RGBColor(226, 239, 251)
ORANGE = RGBColor(232, 126, 45)
RED = RGBColor(190, 50, 50)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(31, 35, 40)
GRAY = RGBColor(96, 105, 115)
LIGHT_GRAY = RGBColor(243, 245, 247)
MID_GRAY = RGBColor(215, 220, 225)
CODE_BG = RGBColor(28, 32, 38)
CODE_OLD = RGBColor(255, 232, 232)
CODE_NEW = RGBColor(231, 247, 224)


def add_text(slide, x, y, w, h, text, size=18, color=BLACK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    p = frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, items, size=17, color=BLACK,
                level_indents=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    for idx, item in enumerate(items):
        if isinstance(item, tuple):
            level, text = item
        else:
            level, text = 0, item
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = ("• " if level == 0 else "  – ") + text
        p.level = level
        p.font.name = "Aptos"
        p.font.size = Pt(size - level)
        p.font.color.rgb = color
        p.space_after = Pt(7 if level == 0 else 3)
        p.line_spacing = 1.05
    return box


def rect(slide, x, y, w, h, fill, line=MID_GRAY, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def title(slide, text, subtitle=None, number=None):
    rect(slide, 0, 0, 13.333, 0.72, NAVY, NAVY)
    add_text(slide, 0.42, 0.13, 11.8, 0.45, text, 26, WHITE, True)
    if subtitle:
        add_text(slide, 0.46, 0.78, 12.1, 0.36, subtitle, 12, GRAY)
    if number is not None:
        add_text(slide, 12.45, 0.16, 0.45, 0.3, str(number), 12, WHITE,
                 True, align=PP_ALIGN.RIGHT)


def footer(slide, source):
    slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0.35), Inches(7.15),
        Inches(12.95), Inches(7.15)
    ).line.color.rgb = MID_GRAY
    add_text(slide, 0.42, 7.18, 12.2, 0.2, source, 8, GRAY)


def code_box(slide, x, y, w, h, label, code, header_color, body_color=CODE_BG,
             font_size=9.2):
    rect(slide, x, y, w, h, body_color, header_color, radius=True)
    rect(slide, x, y, w, 0.4, header_color, header_color)
    add_text(slide, x + 0.12, y + 0.07, w - 0.24, 0.25, label, 13, WHITE, True)
    add_text(slide, x + 0.12, y + 0.48, w - 0.24, h - 0.57, code,
             font_size, WHITE, False, "Menlo")


def benefit_box(slide, y, why, benefit, caveat=None):
    rect(slide, 0.55, y, 12.2, 1.05, LIGHT_GRAY, MID_GRAY, radius=True)
    add_text(slide, 0.78, y + 0.13, 0.72, 0.25, "WHY", 11, DARK_GREEN, True)
    add_text(slide, 1.45, y + 0.1, 5.0, 0.38, why, 12, BLACK)
    add_text(slide, 6.65, y + 0.13, 0.92, 0.25, "BENEFIT", 11, BLUE, True)
    add_text(slide, 7.52, y + 0.1, 4.75, 0.38, benefit, 12, BLACK)
    if caveat:
        add_text(slide, 0.78, y + 0.59, 11.5, 0.27,
                 "Guardrail: " + caveat, 10, RED)


def metric(slide, x, y, w, value, label, fill=LIGHT_BLUE, value_color=BLUE):
    rect(slide, x, y, w, 1.15, fill, fill, radius=True)
    add_text(slide, x + 0.08, y + 0.12, w - 0.16, 0.48, value, 27,
             value_color, True, align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.12, y + 0.7, w - 0.24, 0.28, label, 10, GRAY,
             align=PP_ALIGN.CENTER)


def flow_box(slide, x, y, w, h, text, fill=LIGHT_BLUE, color=NAVY):
    rect(slide, x, y, w, h, fill, color, radius=True)
    add_text(slide, x + 0.08, y + 0.12, w - 0.16, h - 0.2, text, 12,
             color, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def arrow(slide, x1, y1, x2, y2, color=GRAY):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[6]


# 1 — Title
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, NAVY, NAVY)
rect(slide, 0, 5.95, 13.333, 1.55, GREEN, GREEN)
add_text(slide, 0.72, 0.78, 11.8, 0.55,
         "uspGetFinalEffectivePercentage", 34, WHITE, True)
add_text(slide, 0.72, 1.43, 11.7, 0.75,
         "Production (SDT) vs outputV3 Optimization Review", 27, WHITE, True)
add_text(slide, 0.76, 2.45, 10.9, 0.65,
         "Code-level comparison • Spark execution design • measured benefit",
         18, RGBColor(205, 218, 232))
metric(slide, 0.78, 3.6, 2.65, "169.1s", "LATEST PRODUCTION", WHITE, NAVY)
metric(slide, 3.72, 3.6, 2.65, "57.736s", "LATEST OUTPUTV3", WHITE, DARK_GREEN)
metric(slide, 6.66, 3.6, 2.65, "111.364s", "WALL TIME SAVED", WHITE, BLUE)
metric(slide, 9.60, 3.6, 2.65, "65.9%", "FASTER", WHITE, ORANGE)
add_text(slide, 0.78, 6.35, 11.7, 0.4,
         "Prepared from the production source, optimized source, benchmark logs, and the full optimization discussion",
         14, NAVY, True, align=PP_ALIGN.CENTER)
add_text(slide, 0.78, 6.85, 11.7, 0.25,
         "September 24, 2026", 11, NAVY, align=PP_ALIGN.CENTER)


# 2 — Executive summary
slide = prs.slides.add_slide(blank)
title(slide, "Executive summary", "What changed and what it delivered", 2)
metric(slide, 0.55, 1.35, 2.8, "2.93×", "PRODUCTION / OUTPUTV3")
metric(slide, 3.55, 1.35, 2.8, "10.244s", "GAIN FROM 67.980s")
metric(slide, 6.55, 1.35, 2.8, "2.792s", "DATED HELPER (WAS 11.713s)")
metric(slide, 9.55, 1.35, 2.8, "7.736s", "REMAINING TO 50s")
add_bullets(slide, 0.65, 2.9, 5.8, 3.7, [
    "Production semantics remain authoritative; outputV3 changes control flow and safe lineage boundaries.",
    "Independent common reads, branches, modes, checkpoints, effective paths, output builds, and writes now overlap.",
    "The largest measured win came from materializing minimum-quarter and filtered-dated inputs before effective calculation.",
    "Checkpoint mode 4 and 32 shuffle partitions outperformed deferred mode 5 and shuffle 8.",
], 16)
add_bullets(slide, 6.8, 2.9, 5.85, 3.7, [
    "Exact parity is required across FinalEffectivePercentages, FNFinalEffectivePercentages, and SM_FinalEffectivePercentages.",
    "The latest run did not load two shared-helper changes: fn_alloc_pfic_m2 and all_ent_post_tag_m0 were absent.",
    "Therefore, 57.736s reflects the pipeline optimizations, but not every implemented helper optimization.",
    "Current dominant work: CPBT, mode preparation, state lines, PFIC input lines, and output-save overhead.",
], 16)
footer(slide, "Sources: benchmark runs on RunID 17376; outputV3 timing summary; exact-parity benchmark contract")


# 3 — SP architecture
slide = prs.slides.add_slide(blank)
title(slide, "What the SP does", "Business stages and result contract", 3)
labels = [
    ("Common\nconfiguration &\ndimensions", 0.45, 1.45, 1.6),
    ("LT / no-LT\nbase chains", 2.35, 1.45, 1.45),
    ("Mode 1/2/3\npreparation", 4.1, 1.45, 1.45),
    ("Fused CPBT\npriority matching", 5.85, 1.45, 1.55),
    ("Minimum\nquarter", 7.7, 1.45, 1.25),
    ("Dated +\nnon-dated\neffective", 9.25, 1.45, 1.45),
    ("Build + write\n3 outputs", 11.05, 1.45, 1.55),
]
for idx, (text, x, y, w) in enumerate(labels):
    fill = CODE_NEW if idx in (1, 2, 5, 6) else LIGHT_BLUE
    flow_box(slide, x, y, w, 1.05, text, fill)
    if idx < len(labels) - 1:
        arrow(slide, x + w, 1.98, labels[idx + 1][1], 1.98)
add_bullets(slide, 0.65, 3.05, 5.8, 3.45, [
    "Mode 1: look-through allocation path.",
    "Mode 2: PFIC and other footnote augmentation.",
    "Mode 3: state-allocation augmentation.",
    "Mode 4: separate production path because of 704(c) behavior.",
], 17)
add_bullets(slide, 6.75, 3.05, 5.85, 3.45, [
    "Fusion uses _mode on joins, anti-joins, groups, ranks, and outputs.",
    "TrackingKey, Tag, LineTypeID, quarter, transfer flags, and partner keys determine exact matching.",
    "Output schemas and RunID-scoped persistence are part of the SP contract.",
    "Small final row count does not imply a small intermediate Spark plan.",
], 17)
footer(slide, "Sources: output/orchestrator.py; outputV3/stages.py; outputV3/ARCHITECTURE.md")


# 4 — Baseline control flow
slide = prs.slides.add_slide(blank)
title(slide, "Baseline SDT execution shape", "Correct fused business logic, but serial orchestration around it", 4)
code_box(slide, 0.55, 1.25, 5.95, 4.65, "PRODUCTION: SERIAL MODE PREPARATION",
"""for mode in modes_123:
    cfg["mode"] = mode
    cfg["_current_mode"] = mode

    alloc_input = build_allocation_input(...) if mode == 2 else None
    sm_input = build_sm_lookthrough_allocation_input(...) if mode == 3 else None
    lt_input = build_lookthrough_allocation_input(...) if mode == 1 else None

    # mode-specific augmentation and eager checkpoints
    per_mode_data[mode] = {...}""", RED, font_size=9.0)
add_bullets(slide, 6.8, 1.35, 5.8, 4.3, [
    "Production already fuses CPBT and effective business logic across modes using _mode.",
    "However, Pass A iterates modes 1, 2, and 3 serially.",
    "Dated effective is completed before non-dated effective begins.",
    "Per-mode final output construction is serial.",
    "Independent checkpoint actions are paid one after another.",
    "One shared mutable cfg object carries mode-specific state.",
], 16)
benefit_box(slide, 5.95,
            "Establish a fair baseline: preserve fused semantics but identify independent actions around the fused core.",
            "The optimization opportunity is wall-clock overlap and repeated-lineage reduction—not business-rule removal.")
footer(slide, "Production snippet: output/orchestrator.py:827–999")


# 5 — Measurement method
slide = prs.slides.add_slide(blank)
title(slide, "How performance was measured", "Side-by-side execution with exact output reconciliation", 5)
steps = [
    ("1", "Snapshot current\nRunID partitions"),
    ("2", "Run production\nfirst"),
    ("3", "Capture outputs,\nwall time, logs"),
    ("4", "Run outputV3\nwith same inputs"),
    ("5", "Compare all three\nexact fingerprints"),
    ("6", "Restore output\npartitions"),
]
for i, (num, text) in enumerate(steps):
    x = 0.48 + i * 2.12
    rect(slide, x, 1.55, 1.82, 1.25, LIGHT_BLUE, BLUE, radius=True)
    add_text(slide, x + 0.12, 1.68, 0.35, 0.34, num, 18, BLUE, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.43, 1.67, 1.25, 0.72, text, 12, NAVY, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    if i < len(steps) - 1:
        arrow(slide, x + 1.82, 2.18, x + 2.1, 2.18)
add_bullets(slide, 0.65, 3.3, 5.8, 2.85, [
    "Wall time is the primary metric.",
    "Parallel task durations are not summed.",
    "Checkpoint logs identify where lazy plans actually execute.",
    "Stage totals can double-count overlapping work.",
], 17)
add_bullets(slide, 6.75, 3.3, 5.9, 2.85, [
    "Internal cost-sum validation is necessary but not sufficient.",
    "Equal row counts do not prove equal values.",
    "All three tables must match schema, count, and fingerprint.",
    "A failed or non-parity run cannot be promoted.",
], 17)
footer(slide, "Source: outputV3/notebook/benchmark_final_effective_percentage.py")


# 6 — Runtime journey
slide = prs.slides.add_slide(blank)
title(slide, "Runtime journey during today’s optimization", "Measured wall time in seconds (lower is better)", 6)
runs = [
    ("Production\n(latest)", 169.1, NAVY),
    ("Early\noutputV3", 74.5, BLUE),
    ("Mode 5 /\nshuffle 8", 86.554, RED),
    ("Mode 4 /\nshuffle 32", 71.185, BLUE),
    ("CPBT inputs +\nparallel effective", 67.980, BLUE),
    ("Parallel\nboundaries", 64.412, BLUE),
    ("Effective input\nmaterialization", 57.736, GREEN),
]
max_val = 175.0
for i, (label, value, color) in enumerate(runs):
    y = 1.32 + i * 0.72
    add_text(slide, 0.45, y + 0.03, 2.0, 0.42, label, 10, BLACK, True,
             align=PP_ALIGN.RIGHT)
    width = 8.7 * value / max_val
    rect(slide, 2.62, y, width, 0.46, color, color, radius=True)
    add_text(slide, 2.74 + width, y + 0.05, 1.15, 0.28,
             f"{value:.3f}s" if value < 100 else f"{value:.1f}s",
             11, color, True)
add_text(slide, 10.95, 1.38, 1.8, 1.0, "65.9%\nfaster", 28, DARK_GREEN,
         True, align=PP_ALIGN.CENTER)
add_text(slide, 10.95, 2.58, 1.8, 1.1, "111.364s\nsaved", 22, BLUE, True,
         align=PP_ALIGN.CENTER)
add_text(slide, 10.95, 4.1, 1.8, 1.2, "2.93×\nthroughput\nimprovement",
         20, ORANGE, True, align=PP_ALIGN.CENTER)
footer(slide, "Run history reconstructed from today’s benchmark discussion; production varies by pass and cluster state")


# 7 — Scheduler
slide = prs.slides.add_slide(blank)
title(slide, "Change 1 — bounded parallel scheduler", "A controlled concurrency primitive replaces ad-hoc serial execution", 7)
code_box(slide, 0.55, 1.25, 5.95, 4.65, "OLD: IMPLICIT SERIAL CONTROL FLOW",
"""for mode in modes_123:
    prepare_mode(mode)

for mode in valid_modes_123:
    build_final_output(mode)

for target in output_tables:
    write(target)""", RED, font_size=10)
code_box(slide, 6.82, 1.25, 5.95, 4.65, "NEW: NAMED, BOUNDED WAVES",
"""prepared = run_group(
    "mode_prep",
    [(f"mode_{m}", prepare_mode, (m,), {})
     for m in modes_123],
)

assembled = run_group("output_build", [...])
run_group("output_writes", [...])""", GREEN, font_size=10)
benefit_box(slide, 5.95,
            "Modes, branches, and writes are independent after their shared dependencies are materialized.",
            "Four-worker bounded execution exposes overlap while recording each wave’s critical duration.",
            "Every future is observed; any failed task fails the complete SP call.")
footer(slide, "New scheduler: outputV3/orchestrator.py:305–432; group usage: outputV3/pipeline.py")


# 8 — Branches and mode prep
slide = prs.slides.add_slide(blank)
title(slide, "Change 2 — parallel branches and mode preparation", "With-LT/no-LT and modes 1/2/3 no longer wait unnecessarily", 8)
code_box(slide, 0.55, 1.25, 5.95, 4.55, "OLD: NUMERIC MODE LOOP",
"""for mode in modes_123:
    cfg["mode"] = mode
    ...
    if mode == 2:
        build_footnote_...
    if mode == 3:
        build_state_...
    per_mode_data[mode] = {...}""", RED, font_size=9.7)
code_box(slide, 6.82, 1.25, 5.95, 4.55, "NEW: ISOLATED TASKS",
"""chains = run_group(
    "lt_nolt_branches", chain_tasks
)

prepared = run_group(
    "mode_prep",
    [(f"mode_{m}", prepare_mode, (m,), {})
     for m in modes_123],
)""", GREEN, font_size=9.7)
benefit_box(slide, 5.86,
            "Mode 2 footnotes and mode 3 state logic are independent once common chains exist.",
            "Latest waves: LT/no-LT 5.230s; mode prep 10.981s. Serial time would be the sum of sibling tasks.",
            "Each task receives a cfg fork; only checkpoint coordination is shared.")
footer(slide, "Old: output/orchestrator.py:827–999; new: outputV3/pipeline.py:320–327, 589–596")


# 9 — state and boundaries
slide = prs.slides.add_slide(blank)
title(slide, "Change 3 — shared state lineage and concurrent boundaries", "Materialize fan-out once, then overlap independent eager actions", 9)
code_box(slide, 0.55, 1.25, 5.95, 4.65, "OLD: STATE CONSUMED DIRECTLY + SERIAL CHECKPOINTS",
"""all_u, state_lines, state_amounts = (
    build_state_allocation_input(...)
)
non_dated, dated = build_state_entities(
    state_lines, non_dated, dated
)

non_dated = _checkpoint(...nde_pre_cpbt...)
dated = _checkpoint(...de_pre_cpbt...)
transfers = _checkpoint(...txfr_pre_cpbt...)""", RED, font_size=8.9)
code_box(slide, 6.82, 1.25, 5.95, 4.65, "NEW: ONE SHARED SEAM + PARALLEL BOUNDARIES",
"""state_lines = checkpoint(
    spark, state_lines, f"state_lines_m{mode}", cfg
)
non_dated, dated = build_state_entities(...)

mode_boundaries = run_group(
    "mode_prep_boundaries",
    [non_dated_task, dated_task, transfer_task],
)""", GREEN, font_size=8.9)
benefit_box(slide, 5.95,
            "Both state entity branches consume the same four-pass input; pre-CPBT outputs are mutually independent.",
            "Mode 3 fell from roughly 12.35s to about 8s in earlier runs; mode-2 boundary wall became the longest task, not the sum.",
            "state_lines_m3 remains expensive (4.170s) but removing it previously increased downstream replay.")
footer(slide, "New: outputV3/pipeline.py:475–565")


# 10 — CPBT barriers
slide = prs.slides.add_slide(blank)
title(slide, "Change 4 — CPBT input and output barriers", "Keep the expensive matching core isolated from upstream mode lineage", 10)
code_box(slide, 0.55, 1.25, 5.95, 4.65, "OLD: HELPER THEN SEQUENTIAL OUTPUT CHECKPOINTS",
"""fused_temp, fused_transfers = (
    build_cost_percentage_by_type(...)
)
fused_temp = _checkpoint(
    spark, fused_temp, "tcp_by_type_fused", cfg
)
fused_transfers = _checkpoint(
    spark, fused_transfers, "txfr_adj_fused", cfg
)""", RED, font_size=9.1)
code_box(slide, 6.82, 1.25, 5.95, 4.65, "NEW: INPUT BREAKS + PARALLEL OUTPUTS",
"""cpbt_inputs = run_group(
    "cpbt_boundaries",
    [input_non_dated, input_dated],
)
fused_temp, fused_transfers = build_cost_percentage_by_type(...)

cpbt_outputs = run_group(
    "cpbt_boundaries",
    [temp_output, transfer_output],
)""", GREEN, font_size=9.1)
benefit_box(slide, 5.95,
            "The CPBT helper must not replay footnote/state preparation; temp and transfer outputs are independent.",
            "CPBT helper improved 14.222s → 12.328s in the first promoted combination. Latest output wave: 3.498s vs 4.154s summed.",
            "Internal parent/tag fan-out checkpoints are retained; broad checkpoint removal regressed badly.")
footer(slide, "Old: output/orchestrator.py:1033–1052; new: outputV3/pipeline.py:656–726")


# 11 — Missing/final
slide = prs.slides.add_slide(blank)
title(slide, "Change 5 — remove no-op missing work and parallelize final CPBT outputs", "Three independent materializations share the same completed CPBT result", 11)
code_box(slide, 0.55, 1.25, 5.95, 4.65, "OLD: SERIAL POST-CPBT MATERIALIZATION",
"""non_dated, dated = compute_missing_entities(...)
non_dated = _checkpoint(... "nde_post_miss_fused")
dated = _checkpoint(... "de_post_miss_fused")

final_cost = build_final_cost_percentage(
    fused_temp, entity_partners
)
final_cost = _checkpoint(... "final_cost_pct_fused")""", RED, font_size=8.9)
code_box(slide, 6.82, 1.25, 5.95, 4.65, "NEW: IDENTITY + BROADCAST + THREE-WAY WAVE",
"""if cfg["_output_v3_missing_entity_identity"]:
    fused_non_dated, fused_dated = tagged_non_dated, tagged_dated

final_cost = build_final_cost_percentage(
    fused_temp, F.broadcast(entity_partners)
)
outputs = run_group(
    "cpbt_boundaries",
    [post_nd, post_dated, final_cost_task],
)""", GREEN, font_size=8.7)
benefit_box(slide, 5.95,
            "The documented missing-entity path is identity for this fused contract; partner dimension is small.",
            "Latest three-way wave completed in 0.896s. The individual identity gain is small but parity-safe.",
            "The post-missing checkpoints remain, preserving downstream lineage breaks.")
footer(slide, "New: outputV3/pipeline.py:727–791")


# 12 — Effective
slide = prs.slides.add_slide(blank)
title(slide, "Change 6 — materialized effective inputs and parallel effective branches", "The largest measured optimization", 12)
code_box(slide, 0.55, 1.25, 5.95, 4.75, "OLD: REPEATED LAZY INPUTS + SERIAL EFFECTIVE",
"""_, min_q, dated = compute_minimum_quarter(...)

eff_dated, pickup, dated = (
    compute_effective_percentage_dated(
        dated, final_cost, min_q, transfers, ...
    )
)
eff_dated = _checkpoint(... "eff_dt_fused")

eff_nd = compute_effective_percentage_non_dated(...)
eff_nd = _checkpoint(... "eff_nd_fused")""", RED, font_size=8.4)
code_box(slide, 6.82, 1.25, 5.95, 4.75, "NEW: CUT SHARED LINEAGE, THEN OVERLAP",
"""inputs = run_group("effective_inputs", [
    checkpoint(dated, "de_post_minq_fused"),
    checkpoint(min_q, "cost_pct_min_q_fused"),
])

effective = run_group("fused_effective", [
    ("dated", compute_dated_effective, (), {}),
    ("non_dated", compute_non_dated_effective, (), {}),
])""", GREEN, font_size=8.4)
benefit_box(slide, 6.08,
            "Minimum-quarter and filtered-dated plans were replayed by multiple dated actions and both effective branches.",
            "Dated helper: 11.713s → 2.792s. Input wave 1.221s + dated task 3.170s replaced a ~12s critical path.",
            "All proven internal dated checkpoints remain.")
footer(slide, "Old: output/orchestrator.py:1100–1154; new: outputV3/pipeline.py:813–914")


# 13 — Plug/output
slide = prs.slides.add_slide(blank)
title(slide, "Change 7 — parallel plugging, output assembly, and writes", "The final independent work is no longer serialized", 13)
code_box(slide, 0.55, 1.25, 5.95, 4.6, "OLD: SEQUENTIAL CHECKPOINTS AND MODE LOOP",
"""dated, non_dated = apply_plugging(...)
dated = _checkpoint(... "eff_dt_plug_fused")
non_dated = _checkpoint(... "eff_nd_plug_fused")

for mode in valid_modes_123:
    result = build_final_output(...)

_save_results(spark, cfg, statuses)""", RED, font_size=9.0)
code_box(slide, 6.82, 1.25, 5.95, 4.6, "NEW: THREE INDEPENDENT WAVES",
"""plugged = run_group("effective_boundaries", [...])

assembled = run_group(
    "output_build",
    [(f"mode_{m}", assemble, (m,), {}) for m in modes],
)

run_group("output_writes", [
    FinalEffective, FNFinalEffective, SMFinalEffective
])""", GREEN, font_size=8.8)
benefit_box(slide, 5.95,
            "Dated/non-dated plugging and per-mode output filters do not depend on sibling completion; target tables are distinct.",
            "Latest plugging wave 0.745s, output build 1.151s, write wave 2.141s.",
            "Every output future is observed; partial success is not returned.")
footer(slide, "Old: output/orchestrator.py:1150–1216; new: outputV3/pipeline.py:915–987 and outputV3/orchestrator.py")


# 14 — computation rewrites
slide = prs.slides.add_slide(blank)
title(slide, "Change 8 — exact-equivalent computation reductions", "Narrow data and avoid duplicate scans without changing row semantics", 14)
code_box(slide, 0.55, 1.18, 5.95, 4.75, "OLD: FOUR ENTITY SCANS + REDUNDANT DISTINCTS",
"""dated_subset = dated.select(keys).distinct()
nondated_subset = non_dated.select(keys).distinct()

all_entities = (
    dated_subset
    .unionByName(nondated_subset)
    .unionByName(dated_subset.filter(type != cost)
                 .withColumn("TypeID", cost))
    .unionByName(nondated_subset.filter(type != cost)
                 .withColumn("TypeID", cost))
    .distinct()
)""", RED, font_size=8.0)
code_box(slide, 6.82, 1.18, 5.95, 4.75, "NEW: ONE UNION + TYPE EXPANSION + ONE DISTINCT",
"""entity_base = dated.select(keys).unionByName(
    non_dated.select(keys)
)
type_variants = when(
    TypeID.isNotNull() & (TypeID != cost),
    array(TypeID, lit(cost))
).otherwise(array(TypeID))

all_entities = (
    entity_base
    .withColumn("TypeID", explode(type_variants))
    .distinct()
)""", GREEN, font_size=8.0)
benefit_box(slide, 6.03,
            "The original relation is set-equivalent but scans each dated/non-dated source twice.",
            "Reduces scans and shuffles before all_ent_m0; additional changes narrow anti-join keys and prefilter transfer matches.",
            "These shared-helper changes were not active in the latest 57.736s run.")
footer(slide, "Shared helper: output/cost_pct_loader.py; enabled only by _output_v3_* flags")


# 15 — Footnote
slide = prs.slides.add_slide(blank)
title(slide, "Change 9 — Mode 2 PFIC shared lineage and lookup batching", "Target repeated scans inside the 3.982s fn_input_lines_m2 checkpoint", 15)
code_box(slide, 0.55, 1.2, 5.95, 4.75, "OLD: FIVE REPEATED PFIC BASE JOINS",
"""# Passes 2–6 each repeat:
alloc_input.alias("I").join(
    broadcast(PFICFootnoteLineItem).alias("P"),
    (I.LineID == P.LineID)
    & (I.LineTypeID == pfic_lt_id),
    "left",
)

# Six separate scalar first() actions
line_id = table.filter(...).select("LineID").first()""", RED, font_size=8.5)
code_box(slide, 6.82, 1.2, 5.95, 4.75, "NEW: ONE MATERIALIZED BASE + ONE LOOKUP ACTION",
"""shared_pfic_input = (
    alloc_input.alias("I")
    .join(broadcast(pfic_lines).alias("P"), ..., "left")
    .select("I.*", P.LineDescription.alias(...))
)
shared_pfic_input = checkpoint(
    shared_pfic_input, "fn_alloc_pfic_m2"
)

# Six limit(1) branches unioned, then one collect()""", GREEN, font_size=8.3)
benefit_box(slide, 6.08,
            "Passes 2–6 repeat the same allocation/PFIC join; six metadata lookups launch six small Spark actions.",
            "Designed to reduce mode-2 preparation and fn_input_lines_m2 while preserving each pass’s distinct semantics.",
            "Latest run lacked fn_alloc_pfic_m2, proving this code was not loaded; signature guard prevents failure on stale helpers.")
footer(slide, "Shared helper: output/pfic_footnotes.py; guarded call: outputV3/pipeline.py:432–457")


# 16 — Checkpoint strategy
slide = prs.slides.add_slide(blank)
title(slide, "Checkpoint strategy — what worked and what failed", "A lineage barrier can save work or merely move it", 16)
left = [
    "KEEP: mode 4 eager local checkpoints",
    "KEEP: 32 shuffle partitions + AQE",
    "KEEP: state_lines_m3",
    "KEEP: CPBT parent/tag fan-out seams",
    "KEEP: dated internal fan-out seams",
    "KEEP: both CPBT input breaks",
]
right = [
    "REJECTED: mode 5 eager=False",
    "REJECTED: shuffle partitions = 8",
    "REJECTED: broad action-lean checkpoints",
    "REJECTED: broad fused mode prep",
    "REJECTED: candidate-claim CPBT rewrite",
    "REJECTED: mechanical barrier removal",
]
rect(slide, 0.55, 1.25, 5.95, 4.7, CODE_NEW, GREEN, radius=True)
add_text(slide, 0.78, 1.46, 5.45, 0.36, "PROMOTED", 18, DARK_GREEN, True)
add_bullets(slide, 0.78, 1.98, 5.35, 3.7, left, 16)
rect(slide, 6.82, 1.25, 5.95, 4.7, CODE_OLD, RED, radius=True)
add_text(slide, 7.05, 1.46, 5.45, 0.36, "REJECTED / REGRESSED", 18, RED, True)
add_bullets(slide, 7.05, 1.98, 5.35, 3.7, right, 16)
benefit_box(slide, 6.03,
            "Deferred mode 5 hid work in later actions; shuffle 8 removed cluster concurrency.",
            "Restoring mode 4 / shuffle 32 improved 86.554s → 71.185s (15.369s, 17.8%).")
footer(slide, "Evidence: 86.554s regression and 71.185s restored benchmark")


# 17 — stage outcomes
slide = prs.slides.add_slide(blank)
title(slide, "Measured outcome by critical area", "Where the final 57.736 seconds are spent", 17)
headers = ["Area", "Before", "Latest evidence", "Outcome"]
cols = [0.5, 3.2, 5.2, 9.2]
widths = [2.7, 2.0, 4.0, 3.6]
for x, w, text in zip(cols, widths, headers):
    rect(slide, x, 1.2, w, 0.52, NAVY, NAVY)
    add_text(slide, x + 0.08, 1.31, w - 0.16, 0.25, text, 12, WHITE, True)
rows = [
    ("Overall wall", "169.1s prod", "57.736s outputV3", "111.364s / 65.9% faster"),
    ("Dated effective", "11.713s helper", "2.792s helper", "Largest confirmed win"),
    ("Effective wave", "~12s critical", "1.221s inputs + 3.170s", "Shared lazy lineage removed"),
    ("CPBT helper", "12.297s", "12.421s", "Still dominant; helper changes not loaded"),
    ("CPBT output", "3.515s transfer", "3.498s transfer", "No material change yet"),
    ("Mode 2 input", "~4s", "3.982s", "PFIC shared seam not loaded"),
    ("Writes", "~4s save", "2.141s write wave", "Parallel table writes"),
]
for r, row in enumerate(rows):
    y = 1.74 + r * 0.69
    fill = WHITE if r % 2 == 0 else LIGHT_GRAY
    for x, w, text in zip(cols, widths, row):
        rect(slide, x, y, w, 0.66, fill, MID_GRAY)
        c = DARK_GREEN if ("faster" in text or "Largest" in text) else BLACK
        add_text(slide, x + 0.08, y + 0.12, w - 0.16, 0.35, text, 11, c,
                 bold=("faster" in text or "Largest" in text))
footer(slide, "Latest outputV3 process and checkpoint timing: 57.736s wall")


# 18 — Correctness/deployment
slide = prs.slides.add_slide(blank)
title(slide, "Correctness and deployment guardrails", "Performance is valid only when the intended code was loaded and outputs match", 18)
add_bullets(slide, 0.65, 1.28, 5.75, 4.9, [
    "Production output is the source of truth.",
    "Compare all three tables—not only count or cost-sum validation.",
    "Preserve _mode in every fused join, anti-join, grouping, and rank.",
    "Preserve empty-string, null, alias, and duplicate-row semantics.",
    "Keep mode 4 on production control flow.",
    "All shared-helper changes default off without an outputV3 flag.",
], 16)
add_bullets(slide, 6.75, 1.28, 5.85, 4.9, [
    "Databricks can retain stale helper signatures.",
    "The pipeline now inspects checkpoint_fn support before passing it.",
    "Expected checkpoint names prove activation.",
    "Absent fn_alloc_pfic_m2 means PFIC sharing was skipped.",
    "Absent all_ent_post_tag_m0 means the CPBT helper change was skipped.",
    "The latest run contained effective-input checkpoints, so that optimization was active.",
], 16)
rect(slide, 0.75, 6.15, 11.85, 0.64, RGBColor(255, 246, 213), ORANGE, radius=True)
add_text(slide, 0.95, 6.3, 11.45, 0.28,
         "Do not claim the latest timing as parity-approved unless the benchmark’s three exact-comparison rows also PASS.",
         12, RED, True, align=PP_ALIGN.CENTER)
footer(slide, "Deployment evidence is taken from checkpoint names and loaded helper signatures")


# 19 — Next steps
slide = prs.slides.add_slide(blank)
title(slide, "Conclusion and next actions", "What is complete, what remains, and the path below 50 seconds", 19)
metric(slide, 0.62, 1.3, 2.75, "57.736s", "CURRENT BEST")
metric(slide, 3.55, 1.3, 2.75, "7.736s", "GAP TO 50s")
metric(slide, 6.48, 1.3, 2.75, "10.244s", "GAIN FROM 67.980s")
metric(slide, 9.41, 1.3, 2.75, "65.9%", "VS LATEST PRODUCTION")
add_bullets(slide, 0.7, 2.95, 5.8, 3.35, [
    "1. Synchronize and reload shared output helpers.",
    "2. Confirm fn_alloc_pfic_m2 and all_ent_post_tag_m0 appear.",
    "3. Run production → outputV3 exact-parity benchmark.",
    "4. Re-measure fn_input_lines_m2, CPBT helper, and txfr_adj_fused.",
], 16)
add_bullets(slide, 6.75, 2.95, 5.85, 3.35, [
    "5. Keep effective-input materialization—it is the largest verified gain.",
    "6. Optimize only the new measured critical path.",
    "7. Avoid mode 5, shuffle 8, and broad barrier removal.",
    "8. Promote only repeatable, exact-parity improvements.",
], 16)
rect(slide, 0.75, 6.25, 11.85, 0.62, GREEN, GREEN, radius=True)
add_text(slide, 0.95, 6.38, 11.45, 0.3,
         "Result: the optimized SP is nearly three times faster while retaining the production business contract.",
         14, NAVY, True, align=PP_ALIGN.CENTER)
footer(slide, "Detailed reference: SP_OPTIMIZATION_GUIDE.md; reusable workflow: SKILL.md")


prs.core_properties.title = "uspGetFinalEffectivePercentage Optimization Review"
prs.core_properties.subject = "Production SDT versus outputV3 code and runtime comparison"
prs.core_properties.author = "iPACS SDT Optimization"
prs.core_properties.keywords = "Spark, Databricks, stored procedure, outputV3, performance"
prs.save(OUTPUT)
print(OUTPUT)
